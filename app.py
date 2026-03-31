#!/usr/bin/env python3
"""
AWS Security Dashboard — Backend
Flask API serving S3, MFA, and Access Key security checks.
Credentials read from ~/.aws/credentials via AWS CLI.
"""

from flask import Flask, jsonify
from flask_cors import CORS
import subprocess, json, os, base64, csv, io, time
from datetime import datetime, timezone

app = Flask(__name__)
CORS(app)

BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH  = os.path.join(BASE_DIR, "s3_tracker.xlsx")

# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────
def run_aws(cmd):
    r = subprocess.run(cmd, shell=True, capture_output=True, text=True)
    return r.stdout.strip(), r.stderr.strip()


def write_to_live_sheet(all_buckets):
    """Overwrite 'Live Sheet - S3' in s3_tracker.xlsx with fresh scan data."""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        if not os.path.exists(EXCEL_PATH):
            print(f"[Excel] File not found at {EXCEL_PATH} — skipping write.")
            return

        def tb():
            s = Side(style='thin', color="BDC3C7")
            return Border(left=s, right=s, top=s, bottom=s)

        wb = load_workbook(EXCEL_PATH)
        ws = wb["Live Sheet - S3"]

        # Clear rows 4–304
        for row in range(4, 305):
            for col in range(1, 7):
                c = ws.cell(row=row, column=col)
                c.value = ""

        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws["A2"] = (f"🔄 Last Scanned: {ts}   |   🔒 Read-Only — auto-populated on scan"
                    f"   |   {len(all_buckets)} total buckets")

        PUBLIC_RED="C0392B"; PUBLIC_BG="FADBD8"
        PRIV_GREEN="1E8449"; PRIV_BG="D5F5E3"
        NULL_GRAY="7F8C8D";  NULL_BG="F2F3F4"
        BLUE="2980B9"; LR="F4F6F8"; WH="FFFFFF"

        for i, b in enumerate(all_buckets):
            row   = i + 4
            bg    = LR if i % 2 == 0 else WH
            stat  = b.get("status", "Public")
            owner = b.get("owner") or "null"
            det   = b.get("details", {})
            pub   = (stat == "Public")

            def cell(col, val, fc, bg_c, bold=False, italic=False, align="left"):
                c = ws.cell(row=row, column=col, value=val)
                c.font      = Font(name='Arial', size=10, bold=bold, italic=italic, color=fc)
                c.fill      = PatternFill("solid", fgColor=bg_c)
                c.alignment = Alignment(horizontal=align, vertical="center")
                c.border    = tb()

            cell(1, b["bucket"],  BLUE,                     bg,                  align="left")
            cell(2, stat,         PUBLIC_RED if pub else PRIV_GREEN,
                                  PUBLIC_BG  if pub else PRIV_BG,   bold=True, align="center")
            cell(3, owner,        NULL_GRAY if owner=="null" else "2C3E50",
                                  NULL_BG   if owner=="null" else bg,
                                  italic=(owner=="null"), align="center")

            for ci, key in enumerate(["BlockPublicAcls","BlockPublicPolicy","RestrictPublicBuckets"], 4):
                v    = det.get(key)
                lbl  = "✅ Enabled" if v is True else ("❌ Disabled" if v is False else "—")
                fc   = PRIV_GREEN if v is True else (PUBLIC_RED if v is False else NULL_GRAY)
                cbg  = PRIV_BG    if v is True else (PUBLIC_BG  if v is False else NULL_BG)
                cell(ci, lbl, fc, cbg, bold=True, align="center")

        wb.save(EXCEL_PATH)
        print(f"[Excel] Written {len(all_buckets)} rows to Live Sheet.")
    except Exception as e:
        print(f"[Excel write error] {e}")


# ─────────────────────────────────────────────────────────────────────────────
# API — S3
# ─────────────────────────────────────────────────────────────────────────────
@app.route('/api/s3-public')
def s3_public():
    out, err = run_aws("aws s3api list-buckets --query 'Buckets[].Name' --output text")
    if not out and err:
        return jsonify({"error": err}), 500

    buckets     = [b for b in out.split() if b]
    vulnerable  = []
    all_buckets = []

    for bucket in buckets:
        cfg_raw,  _ = run_aws(f"aws s3api get-public-access-block --bucket {bucket} 2>/dev/null")
        tags_raw, _ = run_aws(f"aws s3api get-bucket-tagging --bucket {bucket} 2>/dev/null")

        # Owner tag
        owner = "null"
        if tags_raw:
            try:
                for t in json.loads(tags_raw).get("TagSet", []):
                    if t.get("Key", "").lower() == "owner":
                        owner = t.get("Value", "null")
                        break
            except Exception:
                pass

        if not cfg_raw:
            entry = {
                "bucket": bucket, "status": "Public", "owner": owner,
                "risk": "CRITICAL", "reason": "No Public Access Block configured",
                "details": {"BlockPublicAcls": False, "IgnorePublicAcls": False,
                             "BlockPublicPolicy": False, "RestrictPublicBuckets": False}
            }
            vulnerable.append(entry)
            all_buckets.append(entry)
            continue

        try:
            cfg    = json.loads(cfg_raw).get("PublicAccessBlockConfiguration", {})
            issues = [k for k, v in cfg.items() if not v]
            pub    = len(issues) > 0
            entry  = {
                "bucket": bucket,
                "status": "Public" if pub else "Private",
                "owner":  owner,
                "risk":   "HIGH" if pub else "SAFE",
                "reason": f"{len(issues)} setting(s) disabled" if pub else "All settings enabled",
                "details": cfg
            }
            if pub:
                vulnerable.append(entry)
            all_buckets.append(entry)
        except Exception:
            pass

    write_to_live_sheet(all_buckets)
    return jsonify({
        "count":       len(vulnerable),
        "total":       len(all_buckets),
        "buckets":     vulnerable,
        "all_buckets": all_buckets   # <-- live sheet page uses this
    })


# ─────────────────────────────────────────────────────────────────────────────
# API — S3 All (used by Live Sheet page to get every bucket inc. private)
# ─────────────────────────────────────────────────────────────────────────────
@app.route('/api/s3-all')
def s3_all():
    """Same as s3-public but returns all_buckets (private + public)."""
    result = s3_public()
    # s3_public already returns all_buckets; just unwrap
    data = result.get_json() if hasattr(result, 'get_json') else result[0].get_json()
    return jsonify(data)


# ─────────────────────────────────────────────────────────────────────────────
# API — MFA
# ─────────────────────────────────────────────────────────────────────────────
@app.route('/api/no-mfa')
def no_mfa():
    run_aws("aws iam generate-credential-report")
    time.sleep(4)

    out, err = run_aws("aws iam get-credential-report --query 'Content' --output text")
    if not out and err:
        return jsonify({"error": err}), 500

    decoded = base64.b64decode(out).decode("utf-8")
    reader  = csv.DictReader(io.StringIO(decoded))
    users   = []

    for row in reader:
        if row.get("password_enabled") == "true" and row.get("mfa_active") == "false":
            users.append({
                "username":         row["user"],
                "password_enabled": True,
                "mfa_active":       False,
                "last_login":       row.get("password_last_used", "N/A"),
                "risk":             "HIGH"
            })

    return jsonify({"count": len(users), "users": users})


# ─────────────────────────────────────────────────────────────────────────────
# API — Old Access Keys
# ─────────────────────────────────────────────────────────────────────────────
@app.route('/api/old-keys')
def old_keys():
    users_out, err = run_aws("aws iam list-users --query 'Users[].UserName' --output text")
    if not users_out and err:
        return jsonify({"error": err}), 500

    now  = datetime.now(timezone.utc)
    keys = []

    for username in [u for u in users_out.split() if u]:
        raw, _ = run_aws(
            f"aws iam list-access-keys --user-name {username} "
            f"--query 'AccessKeyMetadata[*].[AccessKeyId,Status,CreateDate]' --output text"
        )
        if not raw:
            continue

        for line in raw.strip().splitlines():
            parts = line.split()
            if len(parts) < 3:
                continue
            key_id, status, created_raw = parts[0], parts[1], parts[2]
            if status != "Active":
                continue
            try:
                dt       = datetime.fromisoformat(
                    created_raw.replace("Z", "+00:00").replace("+00:00+00:00", "+00:00"))
                age_days = (now - dt).days
            except Exception:
                continue

            if age_days > 90:
                keys.append({
                    "username": username, "key_id": key_id,
                    "created":  created_raw, "age_days": age_days,
                    "risk":     "CRITICAL" if age_days > 365 else
                                "HIGH"     if age_days > 180 else "MEDIUM"
                })

    return jsonify({"count": len(keys), "keys": keys})


# ─────────────────────────────────────────────────────────────────────────────
if __name__ == '__main__':
    print("\n✅  AWS Security Dashboard backend running → http://localhost:5001\n")
    app.run(port=5001, debug=False)
